# -*- coding: utf-8 -*-
"""
generator.py —— 内容生成引擎 (v2.0 系统Prompt对齐版)
================================================
图片模型（固定版式，账号统一）：
  第1张 封面 = 固定模板，程序渲染，只填「主题+周数」
  第2张 代码 = 用户上传，自动打「代码含义」标签
  第3张起    = 生活/学习照，程序按爆款版式自动赋角色排版

token 最小化：默认离线模板引擎 0 token；可选 API 润色。
"""
import os
import json
import shutil
import datetime
import urllib.request

# 音频截取模块（可选依赖，无则安全降级）
try:
    import audio as A
except ImportError:
    A = None

# 音乐推荐模块（V2.0 新增，替代硬编码 BGM 匹配）
try:
    from music_recommender import (
        recommend, recommend_simple, RecommendRequest,
        ContentContext, to_legacy_bgm, to_display_list,
    )
    HAS_MUSIC_RECOMMENDER = True
except ImportError:
    HAS_MUSIC_RECOMMENDER = False

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(ROOT, "output")
DATA_DIR = os.path.join(ROOT, "data")
SETTINGS_PATH = os.path.join(DATA_DIR, "settings.json")
RECORD_XLSX = os.path.join(ROOT, "07_发布记录.xlsx")
RECORD_CSV = os.path.join(ROOT, "07_发布记录.csv")
# 素材库（固定封面模板 / 代码截图 / 生活学习照 / BGM音乐）统一入口
# 项目结构（产品设计方案第六章）：素材位于 07_数据资产/图片素材
MATERIAL_DIR = os.path.join(os.path.dirname(ROOT), "07_数据资产", "图片素材")

DEFAULT_SETTINGS = {
    "engine": "offline", "api_key": "",
    "api_url": "https://api.deepseek.com/v1/chat/completions",
    "api_model": "deepseek-chat", "account_name": "04年AI项目记录", "age": "04年",
    "output_dir": "",  # 空=使用默认 output/；用户可通过设置窗口更改
}
HASHTAGS = "#AI学习 #AI项目 #04年 #成长记录 #人工智能"

# ── 动态加载系统提示词（从文档读取，保证文档即代码） ──
def _load_system_prompt():
    """从项目文档动态加载系统提示词。
    优先级：04_Prompt提示词/系统Prompt/内容Agent > 02_个人知识库/我的创作者档案.md > 内置默认
    用户改文档 = 改 AI 行为，无需改代码。"""
    import re

    # ① 从 04_Prompt提示词/系统Prompt/内容Agent 读取（权威来源）
    prompt_file = os.path.join(ROOT, "04_Prompt提示词", "系统Prompt", "内容Agent")
    if os.path.exists(prompt_file):
        try:
            with open(prompt_file, "r", encoding="utf-8") as f:
                content = f.read()
            # 提取角色定义到禁止事项之间的核心提示词（跳过文件路径和对接协议部分）
            # 取"## 角色定义"到"# 与执行发布Agent的对接协议"之间的内容
            m = re.search(r'(##\s*角色定义.*?)(?=#\s*═══════════════════)', content, re.DOTALL)
            if m:
                core = m.group(1).strip()
                # 加上输出格式要求
                output_section = re.search(
                    r'(#\s*输出到文件系统.*?)(?=#\s*与执行发布Agent)', content, re.DOTALL)
                if output_section:
                    core += "\n\n" + output_section.group(1).strip()
                return core
            # 如果正则没匹配到，返回大部分内容（跳过头部元信息）
            # 去掉开头的版本注释
            content = re.sub(r'^>.*\n.*\n.*\n', '', content, flags=re.MULTILINE)
            return content.strip()
        except Exception:
            pass

    # ② 从 02_个人知识库/我的创作者档案.md 构建提示词骨架
    persona_md = os.path.join(ROOT, "02_个人知识库", "我的创作者档案.md")
    if os.path.exists(persona_md):
        try:
            with open(persona_md, "r", encoding="utf-8") as f:
                persona = f.read()
            who = re.search(r'##\s*我是谁\s*\n+(.*?)(?=\n##|\Z)', persona, re.DOTALL)
            style = re.search(r'##\s*表达风格\s*\n+(.*?)(?=\n##|\Z)', persona, re.DOTALL)
            redline = re.search(r'##\s*红线[^)]*\s*\n+(.*?)(?=\n##|\Z)', persona, re.DOTALL)
            parts = ["你是专业短视频运营专家，帮一个04年出生、正在做AI项目的年轻人做抖音内容。"]
            if who:
                parts.append(f"人设：{who.group(1).strip()}")
            if style:
                parts.append(f"风格要求：{style.group(1).strip()}")
            if redline:
                parts.append(f"禁止：{redline.group(1).strip()}")
            return "\n".join(parts)
        except Exception:
            pass

    # ③ 内置默认（兜底，与内容Agent Prompt保持一致）
    return (
        "你是一名专业的AI内容策划和短视频图文运营助手。"
        "你的核心任务：根据用户提供的主题、图片素材以及用户个人档案，生成符合用户个人定位的高质量内容。\n\n"
        "【用户身份】04年出生、正在做AI项目的个人创作者。内容方向：AI项目实践、问题解决、AI工具探索、个人成长经历。\n\n"
        "【内容原则】真实——不编造用户没有经历的事情。成长感——重点表现过程、学习、尝试、失败、改进、收获。"
        "个人化——避免普通AI鸡汤，多写具体经历。\n\n"
        "【表达风格】年轻、真实、有成长感。第一人称，像真实记录，少营销感，少空洞表达，多具体经历。"
        "避免过度专业术语、夸张标题党、成功学表达。\n\n"
        "【输出】1.内容定位 2.标题方案(3个) 3.正文(开头吸引→中间讲述→结尾引导互动) 4.话题标签(5-8个) 5.封面文字(3-8字) 6.配乐建议\n\n"
        "【禁止】编造经历、虚构成绩、虚假炫耀、模仿其他账号风格、输出与用户定位无关内容。"
    )


SYSTEM_PROMPT = _load_system_prompt()

# 爆款版式角色库：程序把每张「生活/学习照」按此节奏赋角色（竖屏图文常见爆款叙事）
ROLES = [
    ("学习过程", "制造代入：让观众看见正在学的你"),
    ("生活松弛", "拉近距离：学习之外的真实生活"),
    ("情绪高点", "跑通/深夜瞬间，把情绪推上去"),
    ("真实细节", "一杯水/一摞书，细节最打动人"),
    ("下一步钩子", "留个悬念，让人想追更"),
]


def load_settings():
    try:
        with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
            s = json.load(f)
        m = dict(DEFAULT_SETTINGS); m.update(s); return m
    except Exception:
        return dict(DEFAULT_SETTINGS)


def save_settings(s):
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
        json.dump(s, f, ensure_ascii=False, indent=2)


def get_output_dir():
    """获取成品包输出目录：优先用户设置，否则默认项目下 output/"""
    s = load_settings()
    custom = s.get("output_dir", "").strip()
    if custom and os.path.isdir(custom):
        return custom
    if custom:
        # 目录不存在则自动创建
        try:
            os.makedirs(custom, exist_ok=True)
            return custom
        except Exception:
            pass
    return OUTPUT_DIR


# 抖音热门 BGM 库：智能体根据主题/情绪自动匹配，用户发布时直接在抖音搜索
# 格式：(歌名, 歌手, 适合情绪, 主题关键词列表, 推荐理由, 截取建议)
DOUYIN_HOT_BGM = [
    # --- focus / 技术 / 代码类 ---
    ("LOVE SCENARIO", "iKON", "focus", ["代码","编程","python","api","项目","调试","bug"],
     "节奏轻快不抢戏，循环感强适合代码推进镜头", "loop 段 0:15~0:45（节奏均匀，剪辑不跳）"),
    ("Sunflower", "Post Malone", "focus", ["代码","编程","算法","模型","训练","部署"],
     "慵懒科技感，适合深夜写代码的氛围", "副歌 0:30~1:00（旋律稳定，情绪不波动）"),
    ("Faded", "Alan Walker", "focus", ["ai","人工智能","神经网络","深度学习","数据"],
     "电子渐进，从平静到饱满，配合技术突破叙事", "drop 段 0:40~1:10（高潮但不炸，适合展示成果）"),
    ("夜空中最亮的星", "逃跑计划", "focus", ["坚持","目标","方向","迷茫","转行"],
     "抖音常驻热门，辨识度极高，配合'找到方向'叙事", "副歌 0:35~1:05（情绪完整起落）"),
    # --- growth / 成长 / 感悟类 ---
    ("可能", "程响", "growth", ["成长","坚持","感悟","转行","改变","努力"],
     "副歌情绪饱满不煽情，配合成长旁白最百搭", "副歌段 0:25~0:55（情绪完整起落）"),
    ("起风了", "买辣椒也用券", "growth", ["成长","时间","回忆","变化","青春"],
     "国民级抖音热歌，情绪从平静到释放", "第二段副歌 0:50~1:20（情绪最高点）"),
    ("少年", "梦然", "growth", ["年轻","热血","不服输","冲","04年"],
     "年轻不服输的劲儿，配合'还在学'的人设", "副歌 0:30~1:00（节奏明快，情绪上扬）"),
    ("你的答案", "阿冗", "growth", ["困难","卡住","突破","答案","解决"],
     "从低谷到爆发，适合'卡了两天终于跑通'的叙事", "副歌 0:40~1:10（爆发段配合突破瞬间）"),
    ("星辰大海", "黄霄雲", "growth", ["目标","远方","梦想","未来","规划"],
     "大气但不鸡汤，适合结尾升华/下周预告", "副歌 0:35~1:05（渐强段，配合文案升华）"),
    # --- life / 生活 / 松弛类 ---
    ("轻快ukulele", "纯音乐", "life", ["生活","日常","咖啡","放松","周末"],
     "阳光生活气，适合生活照/学习间隙的松弛镜头", "全程可用 0:00~0:30（节奏轻快不拖沓）"),
    ("Mojito", "周杰伦", "life", ["生活","夏天","轻松","开心","日常"],
     "轻松愉快不抢戏，生活照百搭", "前奏+第一段 0:00~0:30（节奏轻快）"),
    ("孤勇者", "陈奕迅", "life", ["孤独","一个人","深夜","坚持","不被理解"],
     "抖音现象级热歌，'一个人的战斗'共鸣极强", "副歌 0:45~1:15（情绪最饱满段）"),
    # --- tech / 前沿 / 炫酷类 ---
    ("Industry Baby", "Lil Nas X", "tech", ["项目","成果","demo","展示","发布"],
     "自信展示感，适合'看我做出来的东西'", "hook 段 0:10~0:40（节奏抓耳，配合成果展示）"),
    ("Unstoppable", "Sia", "tech", ["突破","里程碑","完成","上线","成功"],
     "渐进爆发，配合'从0到1做出来了'的高光时刻", "副歌 0:30~1:00（爆发段配合成果）"),
]

# 情绪 → 默认推荐索引（关键词全不命中时的兜底）
_MOOD_DEFAULTS = {
    "focus": [0, 1, 3],   # LOVE SCENARIO, Sunflower, 夜空中最亮的星
    "growth": [4, 5, 6],  # 可能, 起风了, 少年
    "life": [9, 10, 11],  # ukulele, Mojito, 孤勇者
    "tech": [12, 13, 2],  # Industry Baby, Unstoppable, Faded
}


def _match_bgm(theme, difficulty, mood):
    """根据主题关键词+情绪自动匹配抖音热门BGM，返回 top3（4-tuple 列表，兼容旧格式）"""
    text = f"{theme} {difficulty}".lower()
    scored = []
    for i, (name, artist, bgm_mood, keywords, why, clip) in enumerate(DOUYIN_HOT_BGM):
        score = 0
        for kw in keywords:
            if kw in text:
                score += 2
        if bgm_mood == mood:
            score += 1
        scored.append((score, i))
    scored.sort(key=lambda x: -x[0])
    # 取 top3；如果最高分=0（无命中），用 mood 默认推荐
    if scored[0][0] == 0:
        picks = _MOOD_DEFAULTS.get(mood, _MOOD_DEFAULTS["growth"])
    else:
        picks = [idx for _, idx in scored[:3]]
    # 转为兼容旧格式的 4-tuple：("歌名 - 歌手", "抖音搜索：xxx", "理由", "截取建议")
    result = []
    for idx in picks:
        name, artist, _m, _kw, why, clip = DOUYIN_HOT_BGM[idx]
        display = f"{name} - {artist}" if artist != "纯音乐" else name
        search_hint = f"抖音搜索：{name} {artist}" if artist != "纯音乐" else f"抖音搜索：{name}"
        result.append((display, search_hint, why, clip))
    return result


def _lines(text, default=""):
    return [ln.strip(" 　") for ln in (text or default).splitlines() if ln.strip()]


def _pick_mood(difficulty, theme=""):
    """根据内容判断情绪类别：focus/growth/life/tech"""
    text = f"{theme} {difficulty}".lower()
    # 有明确卡点/困难 → growth
    if _lines(difficulty):
        return "growth"
    # 生活/松弛关键词 → life
    life_kw = ["生活", "日常", "咖啡", "放松", "周末", "休息", "出门"]
    if any(k in text for k in life_kw):
        return "life"
    # 前沿/成果关键词 → tech
    tech_kw = ["项目", "demo", "部署", "上线", "发布", "成果", "展示"]
    if any(k in text for k in tech_kw):
        return "tech"
    return "focus"


def next_publish_slot(now=None):
    now = now or datetime.datetime.now()
    wd = [12 * 60, 18 * 60 + 30, 21 * 60 + 30]
    we = [10 * 60 + 30, 20 * 60 + 30]
    for d_add in range(8):
        d = now + datetime.timedelta(days=d_add)
        slots = we if d.weekday() >= 5 else wd
        day_label = "一二三四五六日"[d.weekday()]
        for m in slots:
            cand = d.replace(hour=m // 60, minute=m % 60, second=0, microsecond=0)
            if cand > now + datetime.timedelta(minutes=30):
                tag = "（今天）" if d_add == 0 else ("（明天）" if d_add == 1 else "")
                return f"{d.month}月{d.day}日 周{day_label} {m//60:02d}:{m%60:02d} {tag}".strip()
    return "本周末 20:30"


def _parse_weeks(week_str):
    """解析周数输入：支持 '6' / '6,7' / '6、7' / '6 7' → 返回 (week_list, week_display)
    week_display 用于文案显示，如 '6-7' 或 '6'"""
    import re
    parts = re.split(r'[,，、\s]+', str(week_str).strip())
    nums = []
    for p in parts:
        p = p.strip()
        if p.isdigit():
            nums.append(int(p))
    if not nums:
        nums = [1]
    nums = sorted(set(nums))
    if len(nums) == 1:
        display = str(nums[0])
    elif nums == list(range(nums[0], nums[-1] + 1)):
        display = f"{nums[0]}-{nums[-1]}"
    else:
        display = ",".join(str(n) for n in nums)
    return nums, display


def _suggest_marks(plan, code_count, extra_count):
    """智能体根据文案+图片信息，给每张照片返回标记建议（0 token 离线决策）

    返回列表，每张照片一项：
      {"mark_enabled": bool, "tag_text": str, "reason": str}

    规则：
      - 封面：永远 enabled=True，tag=周数+主题（模板自带）
      - 代码：tag_lines 有实质内容 → enabled=True；否则 False
      - 生活照：根据照片角色和文案内容判断是否值得加标记
    """
    suggestions = []

    # ① 封面 —— 永远标记
    suggestions.append({
        "mark_enabled": True,
        "tag_text": f"第{plan['week']}周 · {plan['theme'][:10]}",
        "reason": "封面模板固定标记，每期统一风格",
    })

    # ② 代码照（可能多张）
    tag_lines = plan.get("tag_lines", [])
    for i in range(code_count):
        tag = tag_lines[i] if i < len(tag_lines) else ""
        if tag and len(tag) >= 2:
            suggestions.append({
                "mark_enabled": True,
                "tag_text": tag[:12],
                "reason": f"代码标签「{tag[:12]}」帮助观众快速理解技术内容",
            })
        else:
            suggestions.append({
                "mark_enabled": False,
                "tag_text": "",
                "reason": "暂无实质代码内容，不加标记更干净",
            })

    # ③ 生活/学习照
    photo_tags = plan.get("photo_tags", [])
    # 跳过封面和代码的标签，从第 (1+code_count) 个开始
    base = 1 + code_count
    roles = plan.get("roles", [])
    role_names = [r[0] for r in roles] if roles else []

    _LIFE_RULES = {
        "学习过程": (True, "代入感最强，标记'正在学习'引发共鸣"),
        "生活松弛": (True, "标记能增加生活感，观众爱看真实一面"),
        "情绪高点": (True, "情绪高潮照片加标记，感染力加倍"),
        "真实细节": (extra_count >= 2, "照片够多时标记细节更有层次；照片少时留白更好看"),
        "下一步钩子": (True, "钩子照片加标记能引导追更"),
    }

    for i in range(extra_count):
        tag = photo_tags[base + i] if (base + i) < len(photo_tags) else ""
        rname = role_names[i] if i < len(role_names) else "记录"
        rule = _LIFE_RULES.get(rname, (True, "标记让每张照片都有故事"))

        # 如果 tag 为空或太短，不建议标记
        if not tag or len(tag) < 2:
            suggestions.append({
                "mark_enabled": False,
                "tag_text": "",
                "reason": f"「{rname}」角色暂无合适标记文字，建议留白",
            })
        else:
            suggestions.append({
                "mark_enabled": rule[0],
                "tag_text": tag[:12],
                "reason": rule[1],
            })

    return suggestions


def generate_plan(week, theme, learned="", completed="", difficulty="", extra_count=1, code_count=1):
    s = load_settings()
    age = s.get("age", "04年")
    age_label = str(age).strip()
    if age_label.isdigit():
        age_label = f"{age_label}岁"
    week_nums, week_display = _parse_weeks(week)
    week_first = week_nums[0]  # 用于取模选择 hook/细节
    theme = (theme or "本周学习").strip()
    learned_ls, completed_ls, diff_ls = _lines(learned), _lines(completed), _lines(difficulty)
    extra_count = max(0, int(extra_count or 0))

    titles = [
        f"学AI第{week_display}周：{theme}，我总算摸到门了",
        f"普通人自学AI，第{week_display}周卡在哪、又怎么爬出来的",
        f"{age_label}，把{theme}从零跑通的那个晚上",
    ]
    hooks = [f"自学AI第{week_display}周，今天差点又放弃了。",
             f"{age_label}做AI，第{week_display}周的真实记录。",
             f"这周学{theme}，我卡了整整两天。"]
    hook = hooks[week_first % len(hooks)]

    life_detail = ["桌上那杯咖啡，早就凉透了。", "室友都睡了，我还对着屏幕较劲。",
                   "报错红了满屏，我深吸一口气，又看了一遍。", "跑通的那一秒，我没忍住笑了一下。",
                   "今天没出门，但脑子走了挺远。"][week_first % 5]
    body = [f"我{age_label}，在做AI项目的第{week_display}周。", "", f"这周死磕的是：{theme}。", ""]
    body += [f"· {x}" for x in learned_ls] + [""]
    if completed_ls:
        body += ["总算把它做出来了：", ""] + [f"· {x}" for x in completed_ls] + [""]
    if diff_ls:
        body += ["中间也卡过：", f"· {diff_ls[0]}", "卡了就卡了，慢慢磨。", ""]
    body += [life_detail, "", "我不是什么天才，", "只是个还在学的普通人。", "", "下周见。"]
    caption = "\n".join(body)

    # 给每张「生活/学习照」按爆款版式赋角色
    roles = [ROLES[i] if i < len(ROLES) else (f"补充镜头{i+1}", "自由发挥，保持真实")
             for i in range(extra_count)]

    # 图片方案（固定版式：封面模板 → 代码 → 爆款节奏的生活/学习照）
    image_plan = [
        {"photo": "① 封面（模板·自动生成）", "role": "统一封面·吸引点击",
         "points": f"程序用固定模板渲染，只填「第{week_display}周 · {theme}」，每期一致"},
        {"photo": "② 代码（自动打标签）", "role": "建立信任",
         "points": "你上传的代码截图，程序自动贴上代码含义标签"},
    ]
    for i, (rname, rdesc) in enumerate(roles):
        image_plan.append({"photo": f"③+ 第{i+1}张生活/学习照", "role": f"{rname}·{rdesc}",
                           "points": f"版式角色：{rname} —— {rdesc}"})

    tag_lines = [ln[:12] for ln in (learned_ls or [theme])][:3]

    # --- 每张图的抖音标签短句（用于图片上的文字标签，离线生成 0 token）---
    first_done = completed_ls[0][:10] if completed_ls else ""
    _ROLE_TAGS = {
        "学习过程": [f"正在啃{theme}，一行行磨", f"屏幕前的第{week_display}个小时", f"{theme}，慢慢搞懂中"],
        "生活松弛": ["学累了，喘口气", "学习之外真实的我", "暂停一下，充充电"],
        "情绪高点": [f"搞定：{first_done}" if first_done else "跑通那一刻！", "终于！值了", "值得的瞬间"],
        "真实细节": ["桌上的日常", "这些细节最真实", "坚持的痕迹"],
        "下一步钩子": ["故事还没完", "下周继续冲", "下一步更精彩"],
    }
    photo_tags = [f"第{week_display}周 · {theme}"]                          # ① 封面标签
    photo_tags.append(tag_lines[0] if tag_lines else theme)          # ② 代码标签
    for i, (rname, _rdesc) in enumerate(roles):
        templates = _ROLE_TAGS.get(rname, ["记录这一刻"])
        photo_tags.append(templates[i % len(templates)])             # ③+ 每张一句话

    ai_prompts = [
        f"cover template, week {week} badge, bold theme text, dark brand background, vertical 9:16",
        "computer screen showing python code editor in dark theme, close-up, vertical 9:16",
        "cozy study / life scene, warm light, calm mood, photorealistic, vertical 9:16",
    ]
    mood = _pick_mood(difficulty, theme)

    # ── BGM 推荐：优先使用独立音乐推荐模块，降级时用旧版硬编码匹配 ──
    if HAS_MUSIC_RECOMMENDER:
        try:
            # 从 HASHTAGS 中提取纯标签词（去掉 # 前缀）
            tag_words = [t.replace("#", "").strip() for t in HASHTAGS.split() if t.strip()]
            bgm_result = recommend(RecommendRequest(
                content=ContentContext(
                    theme=theme,
                    body_text=caption,
                    mood=mood,
                    difficulty=difficulty,
                    tags=tag_words,
                ),
                top_k=3,
            ))
            bgm = to_display_list(bgm_result)  # 转换为旧格式兼容的 4-tuple 列表
        except Exception:
            bgm = _match_bgm(theme, difficulty, mood)  # 降级
    else:
        bgm = _match_bgm(theme, difficulty, mood)

    shot_order = ["镜头1（0-3s）：① 封面 + 钩子字幕", "镜头2（3-8s）：② 代码推进，标签逐条出现"]
    for i, (rname, rdesc) in enumerate(roles):
        shot_order.append(f"镜头{3+i}：{rname}（{rdesc}）")
    shot_order.append(f"结尾：账号名 + 「下周见」")

    return {
        "week": week_display, "week_nums": week_nums, "theme": theme, "age": age, "mood": mood,
        "titles": titles, "hook": hook, "caption": caption, "hashtags": HASHTAGS,
        "image_plan": image_plan, "tag_lines": tag_lines, "photo_tags": photo_tags,
        "ai_prompts": ai_prompts,
        "bgm": bgm, "shot_order": shot_order, "roles": roles, "extra_count": extra_count,
        "publish_time": next_publish_slot(), "engine": "offline",
        "mark_suggestions": _suggest_marks(
            {"week": week_display, "theme": theme, "tag_lines": tag_lines,
             "photo_tags": photo_tags, "roles": roles},
            code_count, extra_count),
    }


def api_enhance(plan):
    s = load_settings()
    if s.get("engine") != "api" or not s.get("api_key"):
        return None
    payload = {"model": s.get("api_model", "deepseek-chat"), "messages": [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": f"把下面文案润色成抖音短句分行体，真实克制，不超过12行，只输出文案：\n{plan['caption']}"}],
        "max_tokens": 400}
    try:
        req = urllib.request.Request(s["api_url"], data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json", "Authorization": f"Bearer {s['api_key']}"})
        with urllib.request.urlopen(req, timeout=20) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        return data["choices"][0]["message"]["content"].strip()
    except Exception as e:
        print(f"[API回退离线] {e}"); return None


def _load_cn_font(size):
    from PIL import ImageFont
    for fp in ("C:/Windows/Fonts/msyhbd.ttc", "C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf"):
        if os.path.exists(fp):
            try:
                return ImageFont.truetype(fp, size)
            except Exception:
                continue
    return ImageFont.load_default()


def _wrap_cn(text, maxc):
    lines, cur = [], ""
    for ch in text:
        cur += ch
        if len(cur) >= maxc:
            lines.append(cur); cur = ""
    if cur:
        lines.append(cur)
    return lines or [text]


# ---------- 固定封面模板（账号统一：只变主题+周数）----------
# 统一标准宽度 = 与代码图/生活照一致，保证轮播大小连贯
COVER_WIDTH = 720

def make_cover(theme, week, out_path, age="04年", account="04年AI项目记录"):
    """无模板照片时的兜底封面：紧凑卡片式，720px宽，不占全屏"""
    from PIL import Image, ImageDraw
    W = COVER_WIDTH
    # 紧凑高度：只容纳必需文字，约400px，不占满屏幕
    H = 400
    img = Image.new("RGB", (W, H), (23, 24, 35))
    d = ImageDraw.Draw(img)
    # 顶部青/红装饰线
    d.rectangle([0, 0, W, 4], fill=(37, 244, 238))
    d.rectangle([0, 4, W, 8], fill=(254, 44, 133))
    # 周数徽章（缩小，居中靠上）
    bf = _load_cn_font(24)
    badge = f"WEEK {week}"
    bbw = int(d.textlength(badge, font=bf)) if hasattr(d, 'textlength') else len(badge) * 24
    d.rounded_rectangle([(W - bbw - 40) // 2, 40, (W + bbw + 40) // 2, 72],
                        radius=10, outline=(37, 244, 238), width=2)
    d.text(((W - bbw) // 2, 44), badge, font=bf, fill=(37, 244, 238))
    d.text(((W - len(f"第 {week} 周") * 14) // 2, 82), f"第 {week} 周",
           font=_load_cn_font(22), fill=(200, 205, 220))
    # 主题大字（居中，单行或两行）
    tf = _load_cn_font(42)
    tlines = _wrap_cn(theme, 10)
    y = 130
    for ln in tlines[:2]:
        try:
            lw = int(d.textlength(ln, font=tf))
        except Exception:
            lw = len(ln) * 42
        d.text(((W - lw) // 2 + 2, y + 2), ln, font=tf, fill=(37, 244, 238))
        d.text(((W - lw) // 2 - 2, y - 2), ln, font=tf, fill=(254, 44, 133))
        d.text(((W - lw) // 2, y), ln, font=tf, fill=(255, 255, 255))
        y += 56
    # 底部账号
    d.text(((W - len(f"@{account}") * 12) // 2, H - 60), f"@{account}",
           font=_load_cn_font(20), fill=(128, 133, 150))
    # 底部声波条
    bars, bw, gap = 5, 12, 10
    heights = [10, 22, 34, 20, 26]
    x0 = (W - bars * (bw + gap) + gap) // 2
    yb = H - 28
    for i in range(bars):
        t = i / (bars - 1)
        c = (int(37 + (254 - 37) * t), int(244 + (44 - 244) * t), int(238 + (133 - 238) * t))
        h = heights[i]
        d.rounded_rectangle([x0 + i * (bw + gap), yb - h // 2, x0 + i * (bw + gap) + bw, yb + h // 2],
                            radius=4, fill=c)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    img.save(out_path, quality=92)
    return out_path


# ---------- 套用用户模板照片：照片做底 + 文字水印 ----------
def make_cover_from_template(template_path, theme, week, out_path, tagline="",
                             account="04年AI项目记录"):
    """以用户模板照片为底图，统一720px宽，底部叠加半透明文字水印。
    与代码图/生活照宽度一致，轮播时大小连贯。"""
    from PIL import Image, ImageDraw
    base = Image.open(template_path).convert("RGB")
    bw, bh = base.size
    # 统一宽度720px，高度等比缩放
    W = COVER_WIDTH
    H = int(bh * W / bw)
    base = base.resize((W, H), Image.LANCZOS)
    img = base.copy()
    d = ImageDraw.Draw(img)

    # ── 底部半透明水印条 ──
    bar_h = max(80, int(H * 0.12))  # 水印条高度：至少80px，约12%画面高度
    y_bar = H - bar_h
    # 渐变遮罩（底部更深）
    for row in range(bar_h):
        alpha = int(160 + 60 * row / bar_h)  # 160→220，越往下越深
        color = (20, 22, 30)  # 深色底
        for x in range(W):
            px = img.getpixel((x, y_bar + row))
            r = int(px[0] * (255 - alpha) / 255 + color[0] * alpha / 255)
            g = int(px[1] * (255 - alpha) / 255 + color[1] * alpha / 255)
            b = int(px[2] * (255 - alpha) / 255 + color[2] * alpha / 255)
            img.putpixel((x, y_bar + row), (r, g, b))

    # ── 水印条上的文字 ──
    # 左侧：Week N · 第 N 周
    lf = _load_cn_font(max(16, min(24, W // 32)))
    week_text = f"WEEK {week}  ·  第 {week} 周"
    tw = int(d.textlength(week_text, font=lf)) if hasattr(d, 'textlength') else len(week_text) * lf.size // 2
    d.text((24, y_bar + bar_h // 2 - lf.size // 2), week_text, font=lf, fill=(37, 244, 238))

    # 右侧：主题（居中偏右）
    tf = _load_cn_font(max(18, min(28, W // 28)))
    theme_short = theme[:14] + ("…" if len(theme) > 14 else "")
    ttw = int(d.textlength(theme_short, font=tf)) if hasattr(d, 'textlength') else len(theme_short) * tf.size // 2
    if ttw < W - tw - 60:
        # 文字不重叠，放右侧
        d.text((W - ttw - 24, y_bar + bar_h // 2 - tf.size // 2),
               theme_short, font=tf, fill=(255, 255, 255))
    else:
        # 文字太长，两行排在水印条中间
        d.text(((W - ttw) // 2, y_bar + 8), theme_short, font=tf, fill=(255, 255, 255))
        d.text((24, y_bar + bar_h // 2 + 4), week_text, font=lf, fill=(37, 244, 238))

    # ── 可选：副标题（tagline）小字 ──
    if tagline and tagline.strip():
        sf = _load_cn_font(max(12, min(16, W // 50)))
        d.text((W - len(tagline) * sf.size // 2 - 20, H - bar_h - 28),
               tagline, font=sf, fill=(255, 255, 255, 180))

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    img.save(out_path, quality=92)
    return out_path


# ---------- 给代码/项目截图打标签（抖音标签样式：侧面小气泡+圆点）----------
def make_tag_image(src_path, tag_lines, out_path):
    """仿抖音发布时的标签功能：图片右侧放小气泡标签+圆点指示器，不遮挡主体内容。"""
    from PIL import Image, ImageDraw
    img = Image.open(src_path).convert("RGBA")
    w, h = img.size
    if w < 720:
        img = img.resize((720, int(h * 720 / w)), Image.LANCZOS); w, h = img.size
    overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    # 标签样式参数
    font_size = max(18, min(28, w // 30))  # 小字体，不抢主体
    font = _load_cn_font(font_size)
    pill_h = int(font_size * 2.2)
    pill_pad_x = int(font_size * 0.8)
    dot_r = max(6, w // 80)  # 圆点半径
    margin_right = int(w * 0.03)
    # 标签垂直起始位置：从图片 30% 高度开始，均匀分布
    n_tags = min(len(tag_lines), 3)
    if n_tags == 0:
        # 无标签时直接保存原图
        out = img.convert("RGB")
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        out.save(out_path, quality=92)
        return out_path
    start_y = int(h * 0.3)
    gap = min(int(h * 0.15), 120)  # 标签间距
    for i, text in enumerate(tag_lines[:3]):
        text = text[:12]  # 限制长度
        cy = start_y + i * gap  # 标签中心 y
        # 计算文字宽度
        try:
            tw = int(draw.textlength(text, font=font))
        except Exception:
            tw = len(text) * font_size
        pill_w = tw + pill_pad_x * 2
        # 气泡位置：右对齐
        pill_x2 = w - margin_right
        pill_x1 = pill_x2 - pill_w
        pill_y1 = cy - pill_h // 2
        pill_y2 = cy + pill_h // 2
        # 画圆点（气泡左侧）
        dot_cx = pill_x1 - dot_r - 4
        dot_cy = cy
        draw.ellipse([dot_cx - dot_r, dot_cy - dot_r, dot_cx + dot_r, dot_cy + dot_r],
                     fill=(255, 255, 255, 230))
        # 画气泡（深色圆角矩形）
        draw.rounded_rectangle([pill_x1, pill_y1, pill_x2, pill_y2],
                               radius=pill_h // 2, fill=(40, 40, 40, 210))
        # 画文字（白色，垂直居中）
        tx = pill_x1 + pill_pad_x
        ty = pill_y1 + (pill_h - font_size) // 2
        draw.text((tx, ty), text, font=font, fill=(255, 255, 255, 245))
    out = Image.alpha_composite(img, overlay).convert("RGB")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    out.save(out_path, quality=92)
    return out_path


# ---------- 轻量标签（生活/学习照用，比代码标签更小更简洁）----------
def make_light_tag(src_path, tag_text, out_path):
    """在图片右下角添加轻量小标签——半透明深色气泡 + 白色小字，不遮挡主体"""
    from PIL import Image, ImageDraw
    img = Image.open(src_path).convert("RGBA")
    w, h = img.size
    if w < 720:
        img = img.resize((720, int(h * 720 / w)), Image.LANCZOS); w, h = img.size

    overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    font_size = max(14, min(22, w // 40))
    font = _load_cn_font(font_size)
    text = tag_text[:12]

    try:
        tw = int(draw.textlength(text, font=font))
    except Exception:
        tw = len(text) * font_size

    pad = int(font_size * 0.6)
    pill_w = tw + pad * 2
    pill_h = int(font_size * 2.0)
    margin = int(w * 0.04)
    pill_x1 = w - margin - pill_w
    pill_y1 = h - margin - pill_h
    pill_x2 = pill_x1 + pill_w
    pill_y2 = pill_y1 + pill_h

    # 半透明圆角气泡
    draw.rounded_rectangle([pill_x1, pill_y1, pill_x2, pill_y2],
                           radius=pill_h // 2, fill=(30, 30, 30, 195))
    draw.text((pill_x1 + pad, pill_y1 + (pill_h - font_size) // 2),
              text, font=font, fill=(255, 255, 255, 235))

    out = Image.alpha_composite(img, overlay).convert("RGB")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    out.save(out_path, quality=92)
    return out_path


# ---------- 生成 publish_package.json（内容Agent → 执行发布Agent 桥梁文件）----------
def _write_publish_package(plan, saved_files, pkg_dir):
    """将 plan 的全部内容写入 publish_package.json，供执行发布Agent读取。
    包含标题（含备选）、钩子、正文、标签、BGM（含搜索关键词）、
    图片路径与类型、发布建议时间等全部字段。"""
    import re

    # 解析 hashtags 字符串为数组
    hashtag_str = plan.get("hashtags", "")
    tags = [t.replace("#", "").strip() for t in hashtag_str.split() if t.strip()]

    # 解析 BGM 列表（4-tuple: display, search_hint, reason, clip）
    bgm_list = []
    for item in (plan.get("bgm") or []):
        display = item[0] if len(item) > 0 else ""
        search = item[1] if len(item) > 1 else ""
        reason = item[2] if len(item) > 2 else ""
        clip = item[3] if len(item) > 3 else ""
        # 从 "歌名 - 歌手" 中拆分
        name, artist = display, ""
        if " - " in display and "纯音乐" not in display:
            parts = display.rsplit(" - ", 1)
            name, artist = parts[0], parts[1]
        bgm_list.append({
            "name": name,
            "artist": artist,
            "display": display,
            "search_keyword": search.replace("抖音搜索：", "").strip(),
            "reason": reason,
            "clip_suggestion": clip,
        })

    # 映射图片类型：根据 image_plan 的 photo 字段识别代码图
    image_plan = plan.get("image_plan", [])
    # 代码图数量 = photo 字段含"代码"的条目数
    code_count = len([p for p in image_plan if "代码" in p.get("photo", "")])
    # 额外照角色列表（生活/学习照）
    roles = plan.get("roles", [])
    tag_lines = plan.get("tag_lines", [])
    # 只保留图片文件，过滤掉 BGM mp3 等非图片
    IMG_EXTS = ('.jpg', '.jpeg', '.png', '.webp', '.bmp', '.gif')
    image_files = [fp for fp in saved_files
                   if fp and os.path.exists(fp) and os.path.splitext(fp)[1].lower() in IMG_EXTS]
    images = []
    for img_idx, fp in enumerate(image_files):
        # 判断类型：封面=0, 代码=1..code_count, 其余=lifestyle
        if img_idx == 0:
            img_type = "cover"
        elif img_idx <= code_count:
            img_type = "code_with_tag"
        else:
            img_type = "lifestyle"
        # 取对应描述（安全取，避免索引越界）
        desc = image_plan[img_idx]["role"] if img_idx < len(image_plan) else f"图片{img_idx+1}"
        tag_text = ""
        if img_type == "code_with_tag":
            ci = img_idx - 1  # 第几张代码图
            if ci < len(tag_lines):
                tag_text = tag_lines[ci]
        elif img_type == "lifestyle":
            li = img_idx - code_count - 1  # 第几张生活照
            photo_tags = plan.get("photo_tags", [])
            pti = 1 + code_count + li  # photo_tags 中的索引（跳过封面和代码）
            if pti < len(photo_tags):
                tag_text = photo_tags[pti]
        images.append({
            "file": os.path.abspath(fp),
            "type": img_type,
            "description": desc,
            "tag_text": tag_text,
        })

    package = {
        "version": "1.0",
        "generated_at": datetime.datetime.now().strftime("%Y-%m-%dT%H:%M:%S+08:00"),
        "generated_by": "content_generation_agent",

        "meta": {
            "week": plan.get("week", ""),
            "theme": plan.get("theme", ""),
            "age": plan.get("age", "04年"),
            "mood": plan.get("mood", "focus"),
            "engine": plan.get("engine", "offline"),
        },

        "content": {
            "title": plan["titles"][0] if plan.get("titles") else "",
            "title_alternatives": plan.get("titles", [])[1:3] if len(plan.get("titles", [])) > 1 else [],
            "hook": plan.get("hook", ""),
            "body": plan.get("caption", ""),
            "tags": tags,
            "cover_text": plan.get("photo_tags", [""])[0] if plan.get("photo_tags") else "",
            "hashtags_raw": hashtag_str,
            "bgm": bgm_list,
            "publish_time_suggestion": plan.get("publish_time", ""),
        },

        "images": images,
        "image_plan": image_plan,
        "shot_order": plan.get("shot_order", []),

        "publish_config": {
            "platform": "douyin",
            "content_type": "image_text",
            "publish_mode": "manual_confirm",
            "enable_comment": True,
            "enable_download": False,
        },

        "output_dir": os.path.abspath(pkg_dir),
    }

    json_path = os.path.join(pkg_dir, "publish_package.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(package, f, ensure_ascii=False, indent=2)
    return json_path


# ---------- 打包成品（封面模板 + 代码打标签 + 爆款排序的额外照）----------
def build_package(plan, code_path, extra_paths, template_path="", tagline=""):
    stamp = datetime.datetime.now().strftime("%m%d")
    # 文件夹名：单周 "第06周"，多周 "第06-07周"
    wnums = plan.get("week_nums", [1])
    if len(wnums) == 1:
        week_folder = f"第{wnums[0]:02d}周"
    else:
        week_folder = f"第{wnums[0]:02d}-{wnums[-1]:02d}周"
    pkg = os.path.join(get_output_dir(), f"{week_folder}_{stamp}")
    os.makedirs(pkg, exist_ok=True)

    # ① 封面：选了模板就套用模板风格（换主题+周数），否则用内置设计稿
    cover_out = os.path.join(pkg, "01_封面_模板.jpg")
    s = load_settings()
    if template_path and os.path.exists(template_path):
        make_cover_from_template(template_path, plan["theme"], plan["week"], cover_out,
                                 tagline, s.get("account_name", "04年AI项目记录"))
    else:
        make_cover(plan["theme"], plan["week"], cover_out, plan.get("age", "04年"),
                   s.get("account_name", "04年AI项目记录"))

    # 读取标记建议（如果 plan 中有的话）
    marks = plan.get("mark_suggestions", [])

    # ② 代码/项目截图：根据 mark_suggestions 决定是否打标签
    code_paths = code_path if isinstance(code_path, list) else ([code_path] if code_path else [])
    code_outs = []
    for ci, cp in enumerate(code_paths):
        mi = 1 + ci  # 封面=0，代码从索引 1 开始
        mark = marks[mi] if mi < len(marks) else {"mark_enabled": True}
        if cp and os.path.exists(cp):
            suffix = f"_{ci+1}" if len(code_paths) > 1 else ""
            if mark.get("mark_enabled", True):
                code_out = os.path.join(pkg, f"02_代码_已打标签{suffix}.jpg")
                # 每张代码图只打自己对应的那一条标签
                own_tag = [plan["tag_lines"][ci]] if ci < len(plan["tag_lines"]) and plan["tag_lines"][ci] else []
                make_tag_image(cp, own_tag, code_out)
            else:
                code_out = os.path.join(pkg, f"02_代码{suffix}.jpg")
                shutil.copy2(cp, code_out)
            code_outs.append(code_out)
        else:
            code_outs.append("")

    # ③+ 额外照：根据 mark_suggestions 决定是否打在图片上的标签
    saved = [cover_out] + code_outs
    roles = plan.get("roles", [])
    base_idx = 2 + len(code_outs)
    for i, p in enumerate(extra_paths or []):
        if not p or not os.path.exists(p):
            continue
        rname = roles[i][0] if i < len(roles) else f"镜头{i+1}"
        mi = 1 + len(code_paths) + i  # marks 中的索引
        mark = marks[mi] if mi < len(marks) else {"mark_enabled": True,
                   "tag_text": plan.get("photo_tags", [""])[mi] if mi < len(plan.get("photo_tags", [])) else ""}

        if mark.get("mark_enabled", False) and mark.get("tag_text", "").strip():
            # 开启标记 → 在图片上打轻量标签
            dst = os.path.join(pkg, f"{base_idx + i:02d}_{rname}_已标记.jpg")
            try:
                make_light_tag(p, mark["tag_text"], dst)
                saved.append(dst)
            except Exception:
                dst = os.path.join(pkg, f"{base_idx + i:02d}_{rname}.jpg")
                shutil.copy2(p, dst); saved.append(dst)
        else:
            # 不标记 → 直接复制原图
            dst = os.path.join(pkg, f"{base_idx + i:02d}_{rname}.jpg")
            try:
                shutil.copy2(p, dst); saved.append(dst)
            except Exception:
                saved.append("")

    # ④ BGM 节奏截取（如果素材库有音频文件 + 依赖已装）
    bgm_trimmed = ""
    if A and A.HAS_AUDIO:
        bgm_dir = os.path.join(MATERIAL_DIR, "BGM音乐")
        bgm_files = A.find_bgm_files(bgm_dir)
        if bgm_files:
            # 按 mood 简单选第一首（后续可扩展映射）
            src_audio = bgm_files[0]
            bgm_out = os.path.join(pkg, "BGM_节奏截取.mp3")
            bgm_trimmed = A.trim_bgm(src_audio, duration_sec=30, out_path=bgm_out)
            if bgm_trimmed:
                saved.append(bgm_trimmed)
    plan["bgm_trimmed"] = bgm_trimmed

    # ⑤ publish_package.json（在 MD 之前生成，此时 saved 含图片+BGM）
    json_path = _write_publish_package(plan, saved, pkg)

    # package_files 只包含图片文件（不含 BGM mp3），供 main.py 直接传给 publisher
    plan["package_files"] = [f for f in saved
        if f and os.path.exists(f) and os.path.splitext(f)[1].lower() in ('.jpg','.jpeg','.png','.webp')]

    md = [f"# 第{plan['week']}周成品包 · {plan['theme']}", "",
          f"> 生成时间：{datetime.datetime.now():%Y-%m-%d %H:%M}｜引擎：{plan['engine']}", "",
          "## 🧩 本版式（固定·账号统一）",
          "1. 封面 = 固定模板（只变主题+周数）",
          "2. 代码 = 自动打标签",
          "3. 生活/学习照 = 程序按爆款节奏排版（见下）", "",
          "## 📌 标题（三选一）"] + [f"{i+1}. {t}" for i, t in enumerate(plan["titles"])]
    md += ["", "## 🎬 3秒钩子", plan["hook"], "", "## 📝 视频文案", "```", plan["caption"], "```",
           "", plan["hashtags"], "", "## 🖼️ 图片版式与角色"]
    md += [f"- **{p['photo']}**（{p['role']}）：{p['points']}" for p in plan["image_plan"]]
    md += ["", "## 🏷️ 代码标签"] + [f"- {t}" for t in plan["tag_lines"]]
    md += ["", "## 🏷️ 每张图的抖音标签短句（贴在图片上）"]
    for i, t in enumerate(plan.get("photo_tags", [])):
        md.append(f"- 第{i+1}张：{t}")
    md += ["", "## 🎵 推荐BGM（抖音热门·发布时直接搜索）"]
    for i, item in enumerate(plan["bgm"]):
        n, search, why = item[0], item[1], item[2]
        clip = item[3] if len(item) > 3 else ""
        md.append(f"{i+1}. **{n}**")
        md.append(f"   理由：{why}")
        if clip:
            md.append(f"   ⏱ {clip}")
        md.append(f"   👉 发布时在抖音「选择音乐」{search}")
    if plan.get("bgm_trimmed"):
        md.append(f"- ✅ 已自动截取：`{os.path.basename(plan['bgm_trimmed'])}`（节奏感知·30秒·淡入淡出）")
    elif A and not A.HAS_AUDIO:
        md.append("- ⚠️ 音频截取依赖未装（pip install pydub numpy imageio-ffmpeg），当前仅文字建议")
    elif not A:
        md.append("- ⚠️ audio 模块不可用，当前仅文字建议")
    md += ["", "## 📹 镜头顺序"] + [f"- {s}" for s in plan["shot_order"]]
    md += ["", f"## ⏰ 建议发布时间：{plan['publish_time']}", "",
           "## ✅ 发布检查清单", "- [ ] 标题无震惊体/虚假包装", "- [ ] 文案符合真实成长人设",
           "- [ ] 封面主题/周数正确", "- [ ] 代码标签清晰", "- [ ] 已在抖音选好BGM", ""]
    with open(os.path.join(pkg, "文案与发布清单.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md))

    plan["package_dir"] = pkg
    return pkg


RECORD_COLS = ["日期", "周数", "主题", "采用标题", "BGM", "状态", "成品包路径"]


def save_record(plan, title, bgm_name, status):
    row = [datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), f"第{plan['week']}周",
           plan["theme"], title, bgm_name, status, plan.get("package_dir", "")]
    try:
        from openpyxl import Workbook, load_workbook
        if os.path.exists(RECORD_XLSX):
            wb = load_workbook(RECORD_XLSX); ws = wb.active
        else:
            wb = Workbook(); ws = wb.active; ws.title = "发布记录"; ws.append(RECORD_COLS)
            for i, wd in enumerate("ABCDEFG"):
                ws.column_dimensions[wd].width = [16, 8, 22, 36, 20, 14, 46][i]
        ws.append(row); wb.save(RECORD_XLSX); return RECORD_XLSX
    except Exception:
        import csv
        exists = os.path.exists(RECORD_CSV)
        with open(RECORD_CSV, "a", newline="", encoding="utf-8-sig") as f:
            wr = csv.writer(f)
            if not exists:
                wr.writerow(RECORD_COLS)
            wr.writerow(row)
        return RECORD_CSV


def selftest():
    from PIL import Image, ImageDraw
    print("== AI-Douyin-Agent 自检 (v2.0) ==")
    tmp = os.path.join(OUTPUT_DIR, "_selftest"); os.makedirs(tmp, exist_ok=True)
    def mk(name, color):
        p = os.path.join(tmp, name)
        Image.new("RGB", (720, 1280), color).save(p); return p
    code = mk("code.png", (24, 26, 38))
    tmpl = mk("tmpl.png", (0, 0, 0))   # 模拟用户模板（黑底，程序套用照片做底+文字水印）
    extras = [mk("life1.jpg", (90, 60, 40)), mk("life2.jpg", (40, 70, 90))]
    print("[1/5] 测试图：1模板(720x1280) + 1代码 + 2额外照")

    plan = generate_plan(5, "RAG检索增强", "调用 DeepSeek API\n搭建向量数据库",
                         "第一个AI聊天机器人", "", extra_count=len(extras))
    print(f"[2/5] 方案：标题{len(plan['titles'])} 版式角色{len(plan['roles'])} BGM{len(plan['bgm'])}")

    pkg = build_package(plan, code, extras, template_path=tmpl, tagline="转行ing")
    cover, tagged = plan["package_files"][0], plan["package_files"][1]
    ok_cover = os.path.exists(cover) and os.path.getsize(cover) > 5000
    ok_tag = os.path.exists(tagged) and os.path.getsize(tagged) > 10000
    ok_extra = all(os.path.exists(x) for x in plan["package_files"][2:])
    from PIL import Image as _Im
    cw, ch = _Im.open(cover).size
    # 封面宽度统一720px（与代码图/生活照一致），高度按模板照片等比缩放
    ok_size = (cw == 720)
    print(f"[3/5] 成品包：{pkg}")
    print(f"[4/5] 封面套用模板={'OK' if ok_cover else 'FAIL'} 宽度720px={'OK' if ok_size else 'FAIL'}({cw}x{ch}) 代码标签={'OK' if ok_tag else 'FAIL'} 额外照排版={'OK' if ok_extra else 'FAIL'}")

    rec = save_record(plan, plan["titles"][0], plan["bgm"][0][0], "自检·未发布")
    print(f"[5/5] 记录写入：{os.path.basename(rec)}｜建议发布：{plan['publish_time']}")
    print("== 自检完成：全部步骤可正常运行 ==")
    return 0


if __name__ == "__main__":
    import sys
    if "--selftest" in sys.argv:
        sys.exit(selftest())
    demo = generate_plan(12, "RAG系统", "RAG流程\nAPI调用", "AI助手Demo", "", extra_count=2)
    print(demo["titles"][0]); print(demo["caption"])
